#!/usr/bin/env python3
"""Gera o resumo da(s) fatura(s): markdown, planilha e PDF.

O markdown é o que vai para a conversa — é o que o usuário lê primeiro. A
planilha é o que vai para o contador e para o arquivo. O PDF é o que se manda
por e-mail para o sócio.

Uso:
    python3 relatorio.py --lancamentos lancamentos.csv [--faturas faturas.csv]
                         [--config config.yaml] [--saida pasta/]
                         [--formatos md,xlsx,pdf] [--titulo "Julho/2026"]
                         [--historico ../2026-06/lancamentos.csv ...]
"""

from __future__ import annotations

import argparse
import collections
import datetime as dt
import html
import os
import shutil
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from comum import (  # noqa: E402
    TIPOS_NAO_DESPESA,
    brl,
    carregar_config,
    despesas,
    ler_faturas,
    ler_lancamentos,
    parse_parcela,
    sem_acento,
)

# Casado com \b nas pontas: sem isso, "max" acha "Obramax" e "prime" acha
# "Primemetal". Recorrência de fornecedor é detectada pelo histórico, não aqui.
ASSINATURA_HINTS = (
    r"netflix|spotify|disney\+?|hbo ?max|globoplay|amazon prime|prime video|"
    r"youtube premium|adobe|microsoft ?365|office ?365|google ?one|icloud|"
    r"dropbox|autodesk|canva|chatgpt|openai|anthropic|claude|sienge|altoqi|"
    r"linkedin|zoom|slack|assinatura|mensalidade"
)

# ---------------------------------------------------------------------------
# Análise
# ---------------------------------------------------------------------------


def agrupar(linhas, chave, rotulo_vazio="(sem)"):
    acc = collections.defaultdict(lambda: {"valor": 0.0, "qtd": 0})
    for l in linhas:
        k = (l.get(chave) or "").strip() or rotulo_vazio
        acc[k]["valor"] = round(acc[k]["valor"] + l["valor"], 2)
        acc[k]["qtd"] += 1
    return sorted(({"chave": k, **v} for k, v in acc.items()), key=lambda x: -x["valor"])


def parcelas_futuras(linhas):
    """Quanto de cada mês seguinte já está comprometido por compras parceladas."""
    futuro = collections.defaultdict(lambda: {"PJ": 0.0, "PF": 0.0, "total": 0.0, "itens": []})
    for l in linhas:
        p = parse_parcela(l.get("parcela"))
        if not p or not l.get("data"):
            continue
        atual, total = p
        try:
            base = dt.date.fromisoformat(l["data"])
        except ValueError:
            continue
        for n in range(1, total - atual + 1):
            mes = (base.month - 1 + n) % 12 + 1
            ano = base.year + (base.month - 1 + n) // 12
            k = f"{ano}-{mes:02d}"
            tit = l.get("titularidade") or "PJ"
            futuro[k][tit if tit in ("PJ", "PF") else "PJ"] += l["valor"]
            futuro[k]["total"] = round(futuro[k]["total"] + l["valor"], 2)
            futuro[k]["itens"].append(
                f"{l.get('estabelecimento') or l.get('descricao')} ({atual + n}/{total})")
    for v in futuro.values():
        v["PJ"], v["PF"] = round(v["PJ"], 2), round(v["PF"], 2)
    return dict(sorted(futuro.items()))


def assinaturas(linhas):
    import re
    rx = re.compile(rf"\b(?:{ASSINATURA_HINTS})\b")
    achadas = collections.defaultdict(lambda: {"valor": 0.0, "qtd": 0, "titularidade": ""})
    for l in linhas:
        if l.get("tipo") == "parcela":
            continue  # compra parcelada não é recorrência, é dívida com prazo
        nome = l.get("estabelecimento") or l.get("descricao")
        if l.get("tipo") == "assinatura" or rx.search(sem_acento(nome)):
            k = nome[:40]
            achadas[k]["valor"] = round(achadas[k]["valor"] + l["valor"], 2)
            achadas[k]["qtd"] += 1
            achadas[k]["titularidade"] = l.get("titularidade") or ""
    return sorted(({"nome": k, **v} for k, v in achadas.items()), key=lambda x: -x["valor"])


def alertas(linhas, faturas, historico_totais, limite_revisao):
    saida = []

    if faturas:
        soma = collections.defaultdict(float)
        for l in linhas:
            soma[l.get("arquivo_origem", "")] += l["valor"]
        for f in faturas:
            dif = round(soma.get(f["arquivo_origem"], 0) - f["total_fatura"], 2)
            if abs(dif) > 0.05:
                saida.append(
                    f"**Fechamento não bate** em `{f['arquivo_origem']}`: extraído "
                    f"{brl(soma.get(f['arquivo_origem'], 0))} contra {brl(f['total_fatura'])} "
                    f"na fatura ({brl(dif)} de diferença). Os números abaixo estão incompletos.")

    vistos = collections.defaultdict(list)
    for l in linhas:
        vistos[(l.get("data"), round(l["valor"], 2), sem_acento(l.get("estabelecimento") or l.get("descricao"))[:25])].append(l)
    for (data, valor, _), grupo in vistos.items():
        if len(grupo) > 1 and valor > 0:
            nome = grupo[0].get("estabelecimento") or grupo[0].get("descricao")
            saida.append(
                f"**Cobrança repetida**: {len(grupo)}x {nome} em {data}, {brl(valor)} cada "
                f"({brl(valor * len(grupo))} no total). Confirme se não é duplicidade do banco.")

    encargos = [l for l in linhas if l.get("tipo") in ("encargo", "anuidade", "iof")]
    if encargos:
        total = round(sum(l["valor"] for l in encargos), 2)
        detalhe = ", ".join(sorted({l.get("tipo", "") for l in encargos}))
        saida.append(f"**Custo do cartão**: {brl(total)} em {detalhe}. "
                     "Vale checar se cabe negociar anuidade ou antecipar pagamento.")

    internacionais = [l for l in linhas if (l.get("moeda_origem") or "").upper() not in ("", "BRL")]
    if internacionais:
        total = round(sum(l["valor"] for l in internacionais), 2)
        saida.append(f"**Compras internacionais**: {len(internacionais)} lançamentos, "
                     f"{brl(total)} já convertidos. Confira se o IOF foi lançado.")

    revisar = [l for l in linhas if l.get("confianca") == "baixa" or not l.get("titularidade")]
    if revisar:
        saida.append(f"**{len(revisar)} lançamento(s) sem classificação segura**, "
                     f"{brl(sum(l['valor'] for l in revisar))} — ver aba/seção de revisão.")

    grandes = [l for l in linhas if l["valor"] >= limite_revisao * 4]
    for l in sorted(grandes, key=lambda x: -x["valor"])[:3]:
        saida.append(f"**Lançamento alto**: {l.get('estabelecimento') or l.get('descricao')} "
                     f"{brl(l['valor'])} em {l.get('data')} "
                     f"({l.get('titularidade') or 'sem classificação'}).")

    if historico_totais:
        anterior = historico_totais[-1]
        atual = round(sum(l["valor"] for l in linhas), 2)
        if anterior["total"]:
            var = (atual - anterior["total"]) / abs(anterior["total"]) * 100
            if abs(var) >= 15:
                direcao = "acima" if var > 0 else "abaixo"
                saida.append(f"**Variação relevante**: total {abs(var):.0f}% {direcao} do mês "
                             f"anterior ({brl(anterior['total'])} → {brl(atual)}).")

    return saida


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def cel(texto, limite=None):
    """Texto seguro para célula de tabela markdown: pipe e quebra de linha fora."""
    t = str(texto or "").replace("|", "/").replace("\n", " ").strip()
    if limite and len(t) > limite:
        t = t[: limite - 1] + "…"
    return t


def gerar_markdown(ctx):
    total = ctx["total"]
    pj, pf = ctx["pj"], ctx["pf"]
    pct = lambda v: f"{(v / total * 100):.0f}%" if total else "—"

    out = [f"# Fatura(s) — {ctx['titulo']}", ""]
    out.append(f"**Total: {brl(total)}**  ·  PJ {ctx['nome_empresa']}: **{brl(pj)}** ({pct(pj)})  "
               f"·  PF pessoal: **{brl(pf)}** ({pct(pf)})")
    if ctx["sem_classificar"]:
        out.append(f"  ·  sem classificação: {brl(ctx['sem_classificar'])}")
    if ctx["pagamentos"]:
        # Sem essa nota, quem confere contra o PDF acha que a conta não bate.
        out.append(f"\nNão entra no total acima: {brl(abs(ctx['pagamentos']))} de pagamento "
                   f"da fatura anterior (aparece no PDF, mas não é gasto do mês).")
    if ctx["historico_totais"]:
        ant = ctx["historico_totais"][-1]
        delta = total - ant["total"]
        sinal = "+" if delta >= 0 else "−"
        out.append(f"\nComparado a {ant['rotulo']}: {sinal}{brl(abs(delta))} "
                   f"({brl(ant['total'])} → {brl(total)}).")
    out.append("")

    if ctx["por_cartao"]:
        out += ["## Por cartão", "", "| Cartão | Lançamentos | Total |", "|---|---:|---:|"]
        for g in ctx["por_cartao"]:
            out.append(f"| {cel(ctx['apelidos'].get(g['chave'], g['chave']))} | {g['qtd']} | {brl(g['valor'])} |")
        out.append("")

    if ctx["por_obra"]:
        out += [f"## PJ por obra / centro de custo", "", "| Centro de custo | Lançamentos | Total | % do PJ |", "|---|---:|---:|---:|"]
        for g in ctx["por_obra"]:
            p = f"{(g['valor'] / pj * 100):.0f}%" if pj else "—"
            out.append(f"| {cel(ctx['nomes_obra'].get(g['chave'], g['chave']))} | {g['qtd']} | {brl(g['valor'])} | {p} |")
        out.append("")

    out += ["## Onde foi o dinheiro", "", "| Categoria | PF/PJ | Total | % |", "|---|:--:|---:|---:|"]
    for g in ctx["por_categoria"][:15]:
        out.append(f"| {cel(g['chave'])} | {ctx['tit_categoria'].get(g['chave'], '—')} | "
                   f"{brl(g['valor'])} | {pct(g['valor'])} |")
    out.append("")

    out += [f"## Maiores lançamentos", "", "| Data | Estabelecimento | Cartão | PF/PJ | Valor |", "|---|---|---|:--:|---:|"]
    for x in ctx["maiores"]:
        out.append(f"| {x.get('data')} | {cel(x.get('estabelecimento') or x.get('descricao'), 38)} | "
                   f"{x.get('cartao')} | {x.get('titularidade') or '?'} | {brl(x['valor'])} |")
    out.append("")

    if ctx["futuras"]:
        out += ["## Já comprometido nos próximos meses", "",
                "Parcelas de compras já feitas — esse valor chega mesmo que ninguém gaste mais nada.",
                "", "| Mês | PJ | PF | Total |", "|---|---:|---:|---:|"]
        for mes, v in list(ctx["futuras"].items())[:12]:
            out.append(f"| {mes} | {brl(v['PJ'])} | {brl(v['PF'])} | {brl(v['total'])} |")
        soma = round(sum(v["total"] for v in ctx["futuras"].values()), 2)
        out += ["", f"**Total ainda a vencer em parcelas: {brl(soma)}**", ""]

    if ctx["assinaturas"]:
        out += ["## Assinaturas e recorrências", "",
                "| Serviço | PF/PJ | Cobranças no mês | No mês | Se repetir 12 meses |",
                "|---|:--:|---:|---:|---:|"]
        for a in ctx["assinaturas"][:15]:
            unitario = a["valor"] / a["qtd"] if a["qtd"] else a["valor"]
            out.append(f"| {cel(a['nome'], 38)} | {a['titularidade'] or '?'} | {a['qtd']} | "
                       f"{brl(a['valor'])} | {brl(unitario * 12)} |")
        out.append("")

    if ctx["alertas"]:
        out += ["## Pontos de atenção", ""]
        out += [f"- {a}" for a in ctx["alertas"]]
        out.append("")

    if ctx["revisar"]:
        out += ["## Pendente de decisão", "",
                "Classifique estes para o próximo mês sair automático:", "",
                "| Estabelecimento | Ocorrências | Valor | Motivo |", "|---|---:|---:|---|"]
        for r in ctx["revisar"][:20]:
            out.append(f"| {cel(r['estabelecimento'], 38)} | {r['ocorrencias']} | {brl(r['valor'])} | {cel(r['motivo'], 70)} |")
        out.append("")

    out += ["---", "",
            "_Separação PF/PJ gerencial, para controle de custo. Não substitui a "
            "análise da contabilidade nem determina dedutibilidade fiscal._"]
    return "\n".join(out)


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

# Ordem das colunas em COLUNAS (comum.py) → letras usadas nas fórmulas
COL = {"valor": "D", "cartao": "G", "tipo": "J", "titularidade": "K",
       "categoria": "L", "centro_custo": "M", "confianca": "N"}
ABA = "'Lançamentos'"


def gerar_xlsx(ctx, caminho):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  ! openpyxl não instalado — planilha não gerada (pip install openpyxl)", file=sys.stderr)
        return None

    from comum import COLUNAS

    wb = Workbook()
    fonte = "Arial"
    titulo_fill = PatternFill("solid", fgColor="1F3864")
    titulo_font = Font(name=fonte, bold=True, color="FFFFFF", size=11)
    negrito = Font(name=fonte, bold=True)
    normal = Font(name=fonte)
    borda = Border(bottom=Side(style="thin", color="BFBFBF"))
    MOEDA = 'R$ #,##0.00;[RED]-R$ #,##0.00'

    def cabecalho(ws, valores):
        for i, v in enumerate(valores, start=1):
            c = ws.cell(row=1, column=i, value=v)
            c.fill, c.font = titulo_fill, titulo_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 28

    # --- Lançamentos ------------------------------------------------------
    ws = wb.active
    ws.title = "Lançamentos"
    cabecalho(ws, [c.capitalize().replace("_", " ") for c in COLUNAS])
    for r, l in enumerate(ctx["linhas"], start=2):
        for i, c in enumerate(COLUNAS, start=1):
            cel = ws.cell(row=r, column=i, value=l.get(c) if c != "valor" else l["valor"])
            cel.font = normal
            cel.border = borda
            if c == "valor":
                cel.number_format = MOEDA
    ult = len(ctx["linhas"]) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{max(ult, 2)}"
    for i, c in enumerate(COLUNAS, start=1):
        larguras = {"descricao": 38, "estabelecimento": 28, "regra": 40, "obs": 24, "categoria": 22}
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(c, 13)

    faixa = lambda col: f"{ABA}!${COL[col]}$2:${COL[col]}${ult}"
    val = f"{ABA}!$D$2:$D${ult}"

    def bloco(ws, linha, titulo, itens, col_chave, extra=(), largura=34):
        """Escreve um bloco 'rótulo | fórmula SUMIFS | fórmula COUNTIFS'.

        Rótulos entre parênteses — '(não alocado)', '(sem cartão)' — são
        placeholders de campo vazio criados por `agrupar`, não valores que
        existem na planilha. Procurar o texto literal daria sempre zero, então
        eles viram critério de célula vazia. `extra` acrescenta condições fixas
        (ex.: só PJ), o que importa justamente no caso do campo vazio.
        """
        ws.cell(row=linha, column=1, value=titulo).font = negrito
        linha += 1
        for cab, c in (("Item", 1), ("Total", 2), ("Qtde", 3)):
            celula = ws.cell(row=linha, column=c, value=cab)
            celula.fill, celula.font = titulo_fill, titulo_font
        linha += 1
        cond_extra = "".join(f',{faixa(c)},"{v}"' for c, v in extra)
        for chave in itens:
            ws.cell(row=linha, column=1, value=chave).font = normal
            texto = str(chave)
            alvo = "" if texto.startswith("(") and texto.endswith(")") else texto.replace('"', '""')
            cond = f'{faixa(col_chave)},"{alvo}"{cond_extra}'
            v = ws.cell(row=linha, column=2, value=f"=SUMIFS({val},{cond})")
            v.number_format, v.font = MOEDA, normal
            q = ws.cell(row=linha, column=3, value=f"=COUNTIFS({cond})")
            q.font = normal
            linha += 1
        ws.column_dimensions["A"].width = largura
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 8
        return linha + 1

    # --- Resumo -----------------------------------------------------------
    ws = wb.create_sheet("Resumo")
    ws.cell(row=1, column=1, value=f"Resumo — {ctx['titulo']}").font = Font(name=fonte, bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Gerado em {dt.date.today().isoformat()}").font = normal
    linha = 4
    for rotulo, formula in (
        ("Total da(s) fatura(s)", f"=SUM({val})"),
        (f"PJ — {ctx['nome_empresa']}", f'=SUMIFS({val},{faixa("titularidade")},"PJ")'),
        ("PF — pessoal", f'=SUMIFS({val},{faixa("titularidade")},"PF")'),
        ("Sem classificação", f'=SUMIFS({val},{faixa("titularidade")},"")'),
        ("A revisar (confiança baixa)", f'=SUMIFS({val},{faixa("confianca")},"baixa")'),
    ):
        ws.cell(row=linha, column=1, value=rotulo).font = negrito
        c = ws.cell(row=linha, column=2, value=formula)
        c.number_format, c.font = MOEDA, negrito
        linha += 1
    linha += 1
    linha = bloco(ws, linha, "Por cartão", [g["chave"] for g in ctx["por_cartao"]], "cartao")
    linha = bloco(ws, linha, "Por tipo de lançamento", [g["chave"] for g in ctx["por_tipo"]], "tipo")

    # --- Por categoria / Por obra ----------------------------------------
    ws = wb.create_sheet("Por categoria")
    bloco(ws, 1, "Gastos por categoria", [g["chave"] for g in ctx["por_categoria"]], "categoria")

    ws = wb.create_sheet("Por obra")
    bloco(ws, 1, "PJ por centro de custo", [g["chave"] for g in ctx["por_obra"]],
          "centro_custo", extra=[("titularidade", "PJ")])

    # --- Parcelas futuras -------------------------------------------------
    ws = wb.create_sheet("Parcelas futuras")
    cabecalho(ws, ["Mês", "PJ", "PF", "Total", "Compras"])
    for r, (mes, v) in enumerate(ctx["futuras"].items(), start=2):
        ws.cell(row=r, column=1, value=mes).font = normal
        for col, chave in ((2, "PJ"), (3, "PF")):
            c = ws.cell(row=r, column=col, value=v[chave])
            c.number_format, c.font = MOEDA, normal
        c = ws.cell(row=r, column=4, value=f"=B{r}+C{r}")
        c.number_format, c.font = MOEDA, negrito
        ws.cell(row=r, column=5, value="; ".join(v["itens"][:8])).font = normal
    for col, w in (("A", 12), ("B", 15), ("C", 15), ("D", 15), ("E", 60)):
        ws.column_dimensions[col].width = w

    # --- Revisar ----------------------------------------------------------
    ws = wb.create_sheet("Revisar")
    cabecalho(ws, ["Data", "Estabelecimento", "Cartão", "Valor", "Sugestão", "Motivo",
                   "→ PF/PJ (preencher)", "→ Centro de custo (preencher)"])
    r = 2
    for l in sorted((x for x in ctx["linhas"]
                     if x.get("confianca") == "baixa" or not x.get("titularidade")),
                    key=lambda x: -abs(x["valor"])):
        ws.cell(row=r, column=1, value=l.get("data")).font = normal
        ws.cell(row=r, column=2, value=(l.get("estabelecimento") or l.get("descricao"))).font = normal
        ws.cell(row=r, column=3, value=l.get("cartao")).font = normal
        c = ws.cell(row=r, column=4, value=l["valor"])
        c.number_format, c.font = MOEDA, normal
        ws.cell(row=r, column=5, value=l.get("titularidade") or "?").font = normal
        ws.cell(row=r, column=6, value=l.get("regra")).font = normal
        for col in (7, 8):
            cel = ws.cell(row=r, column=col, value="")
            cel.fill = PatternFill("solid", fgColor="FFF2CC")  # amarelo = preencher aqui
        r += 1
    ws.cell(row=r + 1, column=1,
            value="Preencha as colunas amarelas e devolva o arquivo: as respostas viram "
                  "regras no config.yaml e não são perguntadas de novo.").font = Font(name=fonte, italic=True)
    for col, w in (("A", 12), ("B", 34), ("C", 9), ("D", 14), ("E", 11), ("F", 44), ("G", 20), ("H", 26)):
        ws.column_dimensions[col].width = w

    wb.save(caminho)
    return caminho


# ---------------------------------------------------------------------------
# PDF (via HTML)
# ---------------------------------------------------------------------------

CSS = """
@page { size: A4; margin: 18mm 14mm; }
body { font-family: Arial, Helvetica, sans-serif; color:#1a1a1a; font-size: 10.5pt; }
h1 { font-size: 18pt; margin: 0 0 2mm; color:#1F3864; }
h2 { font-size: 12.5pt; margin: 7mm 0 2mm; color:#1F3864;
     border-bottom: 1px solid #d0d7e5; padding-bottom: 1mm; }
table { border-collapse: collapse; width: 100%; margin: 2mm 0 4mm; font-size: 9.5pt; }
th { background:#1F3864; color:#fff; text-align:left; padding: 2mm; font-weight:bold; }
td { padding: 1.6mm 2mm; border-bottom: 1px solid #e3e3e3; }
td.num, th.num { text-align: right; }
.destaque { background:#f4f6fb; padding: 3mm 4mm; border-left: 3px solid #1F3864; margin: 3mm 0; }
.rodape { margin-top: 8mm; font-size: 8.5pt; color:#666; font-style: italic;
          border-top: 1px solid #ddd; padding-top: 2mm; }
ul { margin: 2mm 0; padding-left: 5mm; } li { margin-bottom: 1.5mm; }
"""


def md_para_html(md, titulo):
    """Conversor mínimo — o markdown gerado aqui usa só título, tabela, lista e negrito."""
    import re

    linhas, out, tabela = md.split("\n"), [], []

    def fecha_tabela():
        if not tabela:
            return
        corpo = [l for l in tabela if not set(l.replace("|", "").strip()) <= set("-: ")]
        out.append("<table>")
        for i, l in enumerate(corpo):
            cels = [c.strip() for c in l.strip().strip("|").split("|")]
            tag = "th" if i == 0 else "td"
            cls = ['class="num"' if re.match(r"^-?R\$|^\d+$|%$", c) else "" for c in cels]
            out.append("<tr>" + "".join(
                f"<{tag} {k}>{inline(c)}</{tag}>" for c, k in zip(cels, cls)) + "</tr>")
        out.append("</table>")
        tabela.clear()

    def inline(t):
        t = html.escape(t)
        t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
        t = re.sub(r"_(.+?)_", r"<em>\1</em>", t)
        t = re.sub(r"`(.+?)`", r"<code>\1</code>", t)
        return t

    em_lista = False
    for l in linhas:
        if l.startswith("|"):
            tabela.append(l)
            continue
        fecha_tabela()
        if em_lista and not l.startswith("- "):
            out.append("</ul>")
            em_lista = False
        if l.startswith("# "):
            out.append(f"<h1>{inline(l[2:])}</h1>")
        elif l.startswith("## "):
            out.append(f"<h2>{inline(l[3:])}</h2>")
        elif l.startswith("- "):
            if not em_lista:
                out.append("<ul>")
                em_lista = True
            out.append(f"<li>{inline(l[2:])}</li>")
        elif l.strip() == "---":
            continue
        elif l.startswith("_") and l.rstrip().endswith("_"):
            out.append(f'<p class="rodape">{inline(l)}</p>')
        elif l.strip():
            cls = ' class="destaque"' if l.startswith("**Total") else ""
            out.append(f"<p{cls}>{inline(l)}</p>")
    if em_lista:
        out.append("</ul>")
    fecha_tabela()
    return (f"<!doctype html><html lang='pt-BR'><head><meta charset='utf-8'>"
            f"<title>{html.escape(titulo)}</title><style>{CSS}</style></head>"
            f"<body>{''.join(out)}</body></html>")


def gerar_pdf(md, caminho, titulo):
    """Tenta weasyprint, depois Chromium headless. Se nenhum existir, deixa o HTML."""
    caminho_html = caminho.replace(".pdf", ".html")
    with open(caminho_html, "w", encoding="utf-8") as fh:
        fh.write(md_para_html(md, titulo))

    try:
        from weasyprint import HTML  # type: ignore
        HTML(filename=caminho_html).write_pdf(caminho)
        return caminho
    except Exception:
        pass

    for navegador in ("chromium", "chromium-browser", "google-chrome",
                      "/opt/pw-browsers/chromium", "/usr/bin/chromium"):
        exe = shutil.which(navegador) or (navegador if os.path.exists(navegador) else None)
        if not exe:
            continue
        try:
            subprocess.run(
                [exe, "--headless", "--no-sandbox", "--disable-gpu",
                 f"--print-to-pdf={os.path.abspath(caminho)}",
                 "--no-pdf-header-footer", f"file://{os.path.abspath(caminho_html)}"],
                check=True, capture_output=True, timeout=120)
            if os.path.exists(caminho):
                return caminho
        except Exception:
            continue

    print(f"  ! Sem weasyprint ou Chromium: PDF não gerado, HTML pronto em {caminho_html}",
          file=sys.stderr)
    return caminho_html


# ---------------------------------------------------------------------------


def main():
    p = argparse.ArgumentParser(description="Gera o resumo das faturas.")
    p.add_argument("--lancamentos", required=True)
    p.add_argument("--faturas")
    p.add_argument("--config")
    p.add_argument("--saida", default=".")
    p.add_argument("--formatos", default="md,xlsx,pdf")
    p.add_argument("--titulo")
    p.add_argument("--historico", nargs="*", default=[],
                   help="lancamentos.csv de meses anteriores, do mais antigo para o mais novo")
    args = p.parse_args()

    config = carregar_config(args.config)
    todas = ler_lancamentos(args.lancamentos)
    linhas = despesas(todas)  # pagamento da fatura anterior não é gasto do mês
    faturas = ler_faturas(args.faturas) if args.faturas else []
    os.makedirs(args.saida, exist_ok=True)

    datas = sorted(l["data"] for l in linhas if l.get("data"))
    titulo = args.titulo or (f"{datas[0]} a {datas[-1]}" if datas else dt.date.today().isoformat())

    historico_totais = []
    for h in args.historico:
        try:
            ant = despesas(ler_lancamentos(h))
        except OSError:
            continue
        historico_totais.append({
            "rotulo": os.path.basename(os.path.dirname(os.path.abspath(h))) or h,
            "total": round(sum(x["valor"] for x in ant), 2),
        })

    por_categoria = agrupar(linhas, "categoria", "(sem categoria)")
    tit_categoria = {}
    for g in por_categoria:
        tits = {l.get("titularidade") for l in linhas
                if (l.get("categoria") or "(sem categoria)") == g["chave"]}
        tits.discard("")
        tit_categoria[g["chave"]] = "/".join(sorted(tits)) if tits else "—"

    revisar_grupos = collections.defaultdict(lambda: {"ocorrencias": 0, "valor": 0.0, "motivo": ""})
    for l in linhas:
        if l.get("confianca") == "baixa" or not l.get("titularidade"):
            k = (l.get("estabelecimento") or l.get("descricao"))[:38]
            revisar_grupos[k]["ocorrencias"] += 1
            revisar_grupos[k]["valor"] = round(revisar_grupos[k]["valor"] + l["valor"], 2)
            revisar_grupos[k]["motivo"] = l.get("regra", "")

    limite = float((config.get("relatorio") or {}).get("limite_revisao_obrigatoria", 500) or 500)
    top_n = int((config.get("relatorio") or {}).get("top_lancamentos", 10) or 10)

    ctx = {
        "titulo": titulo,
        "linhas": linhas,
        "total": round(sum(l["valor"] for l in linhas), 2),
        "pj": round(sum(l["valor"] for l in linhas if l.get("titularidade") == "PJ"), 2),
        "pf": round(sum(l["valor"] for l in linhas if l.get("titularidade") == "PF"), 2),
        "sem_classificar": round(sum(l["valor"] for l in linhas if not l.get("titularidade")), 2),
        "pagamentos": round(sum(l["valor"] for l in todas
                                if l.get("tipo") in TIPOS_NAO_DESPESA), 2),
        "nome_empresa": (config.get("empresa") or {}).get("apelido")
                        or (config.get("empresa") or {}).get("nome") or "empresa",
        "apelidos": {str(c.get("final", ""))[-4:]: c.get("apelido") or c.get("banco") or str(c.get("final"))
                     for c in (config.get("cartoes") or [])},
        "nomes_obra": {o.get("codigo"): o.get("nome") or o.get("codigo")
                       for o in (config.get("obras") or [])},
        "por_cartao": agrupar(linhas, "cartao", "(sem cartão)"),
        "por_categoria": por_categoria,
        "tit_categoria": tit_categoria,
        "por_tipo": agrupar(linhas, "tipo", "(sem tipo)"),
        "por_obra": agrupar([l for l in linhas if l.get("titularidade") == "PJ"],
                            "centro_custo", "(não alocado)"),
        "maiores": sorted(linhas, key=lambda x: -x["valor"])[:top_n],
        "futuras": parcelas_futuras(linhas),
        "assinaturas": assinaturas(linhas),
        "historico_totais": historico_totais,
        "revisar": sorted(({"estabelecimento": k, **v} for k, v in revisar_grupos.items()),
                          key=lambda x: -abs(x["valor"])),
    }
    ctx["alertas"] = alertas(linhas, faturas, historico_totais, limite)

    formatos = {f.strip() for f in args.formatos.split(",") if f.strip()}
    md = gerar_markdown(ctx)
    gerados = []

    if "md" in formatos:
        caminho = os.path.join(args.saida, "resumo.md")
        with open(caminho, "w", encoding="utf-8") as fh:
            fh.write(md)
        gerados.append(caminho)
    if "xlsx" in formatos:
        r = gerar_xlsx(ctx, os.path.join(args.saida, "resumo.xlsx"))
        if r:
            gerados.append(r)
    if "pdf" in formatos:
        gerados.append(gerar_pdf(md, os.path.join(args.saida, "resumo.pdf"), titulo))

    print(md)
    print("\n---\nArquivos gerados:")
    for g in gerados:
        print(f"  · {g}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
