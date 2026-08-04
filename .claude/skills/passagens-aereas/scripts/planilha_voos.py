#!/usr/bin/env python3
"""Gera a planilha comparativa de voos (.xlsx) a partir de um JSON.

Uso:
    python planilha_voos.py opcoes.json voos.xlsx

Formato do JSON de entrada (só "opcoes" é obrigatório):

{
  "rota": "GRU -> LIS",
  "periodo": "10 a 25 de setembro de 2026",
  "passageiros": "1 adulto, econômica",
  "consultado_em": "2026-08-04",
  "fontes": "Google Flights, Skyscanner",
  "opcoes": [
    {
      "companhia": "TAP",
      "ida": "10/09 22:05",
      "volta": "25/09 11:30",
      "duracao": "10h20",
      "paradas": "direto",
      "bagagem": "1 despachada 23kg incluída",
      "preco_total": 4180.00,
      "moeda": "BRL",
      "link": "https://...",
      "observacao": "voo noturno"
    }
  ],
  "datas_alternativas": [
    {"data": "08/09 (ter)", "preco": 3740.00, "diferenca": -440.00}
  ],
  "recomendacao": "Opção 1: direto, noturno e com bagagem inclusa."
}

A aba principal fica ordenada por preço, com a linha mais barata destacada.
Requer openpyxl (pip install openpyxl).
"""

import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl não está instalado. Rode: pip install openpyxl")

CABECALHO = PatternFill("solid", fgColor="1F3864")
DESTAQUE = PatternFill("solid", fgColor="E2EFDA")
BORDA = Border(bottom=Side(style="thin", color="D9D9D9"))

COLUNAS = [
    ("#", 5),
    ("Companhia", 20),
    ("Ida", 16),
    ("Volta", 16),
    ("Duração", 12),
    ("Paradas", 22),
    ("Bagagem", 32),
    ("Preço total", 14),
    ("Observação", 30),
    ("Conferir", 14),
]


def _formato_moeda(moeda):
    return 'R$ #,##0.00' if (moeda or "BRL").upper() == "BRL" else '#,##0.00'


def _titulo(ws, dados, ncols):
    linhas = []
    rota = dados.get("rota")
    periodo = dados.get("periodo")
    if rota or periodo:
        linhas.append(" · ".join(p for p in (rota, periodo) if p))
    if dados.get("passageiros"):
        linhas.append(dados["passageiros"])

    for i, texto in enumerate(linhas, start=1):
        ws.cell(row=i, column=1, value=texto).font = Font(
            bold=(i == 1), size=14 if i == 1 else 11, color="1F3864"
        )
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=ncols)
    return len(linhas) + 2 if linhas else 1


def _tabela_opcoes(ws, dados, primeira_linha):
    opcoes = sorted(
        dados["opcoes"],
        key=lambda o: o.get("preco_total") if o.get("preco_total") is not None else float("inf"),
    )

    for col, (titulo, largura) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=primeira_linha, column=col, value=titulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = CABECALHO
        celula.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.row_dimensions[primeira_linha].height = 22

    for i, opcao in enumerate(opcoes):
        linha = primeira_linha + 1 + i
        valores = [
            i + 1,
            opcao.get("companhia", ""),
            opcao.get("ida", ""),
            opcao.get("volta", ""),
            opcao.get("duracao", ""),
            opcao.get("paradas", ""),
            opcao.get("bagagem", ""),
            opcao.get("preco_total"),
            opcao.get("observacao", ""),
        ]
        for col, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=col, value=valor)
            celula.border = BORDA
            celula.alignment = Alignment(vertical="center", wrap_text=col in (6, 7, 9))
            if i == 0:
                celula.fill = DESTAQUE
        preco = ws.cell(row=linha, column=8)
        preco.number_format = _formato_moeda(opcao.get("moeda"))
        preco.font = Font(bold=(i == 0))

        link = ws.cell(row=linha, column=10)
        link.border = BORDA
        if i == 0:
            link.fill = DESTAQUE
        if opcao.get("link"):
            link.value = "abrir"
            link.hyperlink = opcao["link"]
            link.font = Font(color="0563C1", underline="single")

    return primeira_linha + len(opcoes) + 2


def _datas_alternativas(ws, dados, linha):
    alternativas = dados.get("datas_alternativas") or []
    if not alternativas:
        return linha

    ws.cell(row=linha, column=1, value="Sair em outras datas").font = Font(bold=True, size=12)
    linha += 1
    for col, titulo in enumerate(["Data", "Preço a partir de", "Diferença"], start=1):
        celula = ws.cell(row=linha, column=col, value=titulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = CABECALHO

    moeda = (dados.get("opcoes") or [{}])[0].get("moeda")
    for alt in alternativas:
        linha += 1
        ws.cell(row=linha, column=1, value=alt.get("data", "")).border = BORDA
        preco = ws.cell(row=linha, column=2, value=alt.get("preco"))
        preco.number_format = _formato_moeda(moeda)
        preco.border = BORDA
        dif = ws.cell(row=linha, column=3, value=alt.get("diferenca"))
        dif.number_format = _formato_moeda(moeda)
        dif.border = BORDA
        if isinstance(alt.get("diferenca"), (int, float)):
            dif.font = Font(color="C00000" if alt["diferenca"] > 0 else "1E7B34")

    return linha + 2


def _rodape(ws, dados, linha, ncols):
    if dados.get("recomendacao"):
        celula = ws.cell(row=linha, column=1, value=f"Recomendação: {dados['recomendacao']}")
        celula.font = Font(bold=True)
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols)
        ws.row_dimensions[linha].height = 30
        linha += 2

    partes = []
    if dados.get("consultado_em"):
        partes.append(f"Preços consultados em {dados['consultado_em']}")
    if dados.get("fontes"):
        partes.append(f"via {dados['fontes']}")
    aviso = " ".join(partes)
    aviso = (aviso + ". " if aviso else "") + (
        "Passagem muda de preço a qualquer momento: confirme o valor final na tela de "
        "pagamento antes de comprar."
    )
    celula = ws.cell(row=linha, column=1, value=aviso)
    celula.font = Font(italic=True, size=9, color="808080")
    celula.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols)
    ws.row_dimensions[linha].height = 28


def gerar(dados, saida):
    if not dados.get("opcoes"):
        sys.exit('O JSON precisa de pelo menos uma entrada em "opcoes".')

    wb = Workbook()
    ws = wb.active
    ws.title = "Voos"
    ncols = len(COLUNAS)

    cabecalho = _titulo(ws, dados, ncols)
    linha = _tabela_opcoes(ws, dados, cabecalho)
    linha = _datas_alternativas(ws, dados, linha)
    _rodape(ws, dados, linha, ncols)

    ws.freeze_panes = f"A{cabecalho + 1}"
    wb.save(saida)
    return saida


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    entrada, saida = Path(sys.argv[1]), Path(sys.argv[2])
    dados = json.loads(entrada.read_text(encoding="utf-8"))
    gerar(dados, saida)
    print(f"Planilha gerada: {saida}")


if __name__ == "__main__":
    main()
